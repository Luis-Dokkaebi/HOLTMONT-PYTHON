"""El pipeline del dashboard de la Pre Work Order (adaptación de `gs3.py`).

`agente_full/pipeline_dashboard.py` corre en el GitHub Action del repositorio
`Agente_Full`: lee el documento que se subió a la Pre Work Order, tres agentes
de Gemini escriben un `app.py` de Streamlit, se valida y queda listo para el
commit que Streamlit Cloud vuelve a desplegar.

Qué es real y qué no en estas pruebas:

* **El modelo de lenguaje es lo único simulado.** Es la frontera con un
  servicio externo de pago; se sustituye por un guion de respuestas.
* La lectura de archivos, AppTest (en su subproceso sin secretos), la descarga
  (contra un servidor HTTP local) y la escritura de `app.py` son las de verdad.
* Las comprobaciones en navegador viven en `test_agente_full_navegador.py`.
"""

from __future__ import annotations

import base64
import http.server
import threading
from importlib import metadata
from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
import pytest
from langchain_core.messages import AIMessage

from agente_full import pipeline_dashboard as pipeline

APP_BUENA = 'import streamlit as st\nst.title("Levantamiento")\nst.write("ok")\n'
APP_QUE_TRUENA = "import streamlit as st\nst.title('x')\nvalor = 1 / 0\n"
JSON_EXTRAIDO = '{"metadata": {"document_title": "Mezanine"}, "tables": []}'


class ModeloGuionado:
    """Un LLM que contesta, en orden, lo que dice su guion y anota qué recibió."""

    def __init__(self, respuestas: List[Any]) -> None:
        self._respuestas = list(respuestas)
        self.recibidos: List[List[Any]] = []

    def invoke(self, mensajes: List[Any]) -> AIMessage:
        self.recibidos.append(mensajes)
        indice = min(len(self.recibidos), len(self._respuestas)) - 1
        return AIMessage(content=self._respuestas[indice])


def _bloque(codigo: str, lenguaje: str = "python") -> str:
    return f"Aquí está:\n```{lenguaje}\n{codigo}\n```\nSaludos"


def _modelos(desarrollador: ModeloGuionado, extractor: ModeloGuionado = None) -> pipeline.Modelos:
    return pipeline.Modelos(
        extractor=extractor or ModeloGuionado([_bloque(JSON_EXTRAIDO, "json")]),
        ux=ModeloGuionado(["Pestañas por tabla y métricas arriba."]),
        desarrollador=desarrollador,
    )


def _navegador_aprueba(codigo: str):
    return True, "ok", None


def _documento(tmp_path: Path) -> Path:
    ruta = tmp_path / "levantamiento.csv"
    ruta.write_text("concepto,monto\nCimentación,1500\n", encoding="utf-8")
    return ruta


# --- Lectura del documento --------------------------------------------------

def test_csv_en_utf8_conserva_los_acentos(tmp_path):
    ruta = tmp_path / "datos.csv"
    ruta.write_text("concepto,monto\nCimentación,1500\n", encoding="utf-8")

    texto = pipeline.leer_archivo(str(ruta))

    assert "| concepto | monto |" in texto
    assert "| Cimentación | 1500 |" in texto


def test_csv_en_latin1_se_lee_igual(tmp_path):
    ruta = tmp_path / "datos.csv"
    ruta.write_bytes("concepto,monto\nCimentación,1500\n".encode("latin-1"))

    assert "| Cimentación | 1500 |" in pipeline.leer_archivo(str(ruta))


def test_las_columnas_y_filas_vacias_no_llegan_al_modelo(tmp_path):
    ruta = tmp_path / "datos.csv"
    ruta.write_text("concepto,vacia,monto\nMuro,,10\n,,\n", encoding="utf-8")

    texto = pipeline.leer_archivo(str(ruta))

    assert "vacia" not in texto
    assert texto.count("\n") == 2  # encabezado, separador y una sola fila


def test_excel_entrega_una_seccion_por_hoja(tmp_path):
    ruta = tmp_path / "junta.xlsx"
    with pd.ExcelWriter(ruta) as libro:
        pd.DataFrame({"material": ["Block"], "cantidad": [450]}).to_excel(libro, sheet_name="Materiales", index=False)
        pd.DataFrame({"categoria": ["Albañil"], "semanas": [2]}).to_excel(libro, sheet_name="Mano de obra", index=False)

    texto = pipeline.leer_archivo(str(ruta))

    assert "### PESTAÑA/HOJA EXCEL: Materiales" in texto
    assert "| Block | 450 |" in texto
    assert "### PESTAÑA/HOJA EXCEL: Mano de obra" in texto
    assert "| Albañil | 2 |" in texto


def _pdf_con_texto(texto: str) -> bytes:
    """Un PDF mínimo de una página con `texto`, armado a mano (sin dependencias)."""
    flujo = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode("latin-1")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(flujo)).encode() + b" >>\nstream\n" + flujo + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    salida = b"%PDF-1.4\n"
    posiciones = []
    for numero, cuerpo in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += f"{numero} 0 obj\n".encode() + cuerpo + b"\nendobj\n"
    inicio_xref = len(salida)
    salida += f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode()
    salida += b"".join(f"{p:010d} 00000 n \n".encode() for p in posiciones)
    salida += f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{inicio_xref}\n%%EOF\n".encode()
    return salida


def test_pdf_entrega_el_texto_por_pagina(tmp_path):
    ruta = tmp_path / "cotizacion.pdf"
    ruta.write_bytes(_pdf_con_texto("TOTAL 1500 MXN"))

    texto = pipeline.leer_archivo(str(ruta))

    assert "--- PÁGINA 1 ---" in texto
    assert "TOTAL 1500 MXN" in texto


def test_word_entrega_parrafos_y_tablas(tmp_path):
    import docx

    ruta = tmp_path / "minuta.docx"
    documento = docx.Document()
    documento.add_paragraph("Minuta de la junta interdisciplinaria")
    tabla = documento.add_table(rows=1, cols=2)
    tabla.rows[0].cells[0].text = "Partida"
    tabla.rows[0].cells[1].text = "Monto"
    documento.save(ruta)

    texto = pipeline.leer_archivo(str(ruta))

    assert "Minuta de la junta interdisciplinaria" in texto
    assert "--- TABLA 1 DE WORD ---" in texto
    assert "Partida | Monto" in texto


def test_la_imagen_llega_como_contenido_multimodal_con_mime_valido(tmp_path):
    ruta = tmp_path / "pizarron.JPG"
    ruta.write_bytes(b"\xff\xd8\xff\xe0contenido")

    contenido = pipeline.leer_archivo(str(ruta))

    assert contenido[0]["type"] == "text"
    url = contenido[1]["image_url"]["url"]
    assert url.startswith("data:image/jpeg;base64,")
    assert base64.b64decode(url.split(",", 1)[1]) == b"\xff\xd8\xff\xe0contenido"


def test_el_texto_plano_se_recorta_a_diez_mil_caracteres(tmp_path):
    ruta = tmp_path / "notas.txt"
    ruta.write_text("x" * 20000, encoding="utf-8")

    assert len(pipeline.leer_archivo(str(ruta))) == 10000


def test_un_archivo_ilegible_detiene_el_pipeline(tmp_path):
    ruta = tmp_path / "roto.xlsx"
    ruta.write_bytes(b"esto no es un excel")

    with pytest.raises(pipeline.ArchivoIlegible, match="roto.xlsx"):
        pipeline.leer_archivo(str(ruta))


def test_un_doc_antiguo_pide_guardarlo_como_docx(tmp_path):
    ruta = tmp_path / "viejo.doc"
    ruta.write_bytes(b"\xd0\xcf\x11\xe0")

    with pytest.raises(pipeline.ArchivoIlegible, match=".docx"):
        pipeline.leer_archivo(str(ruta))


def test_la_tabla_markdown_no_se_rompe_con_barras_saltos_ni_vacios():
    df = pd.DataFrame({"nota": ["a|b\nc", None], "monto": [1.5, float("nan")]})

    lineas = pipeline.tabla_markdown(df).splitlines()

    assert lineas[0] == "| nota | monto |"
    assert lineas[1] == "|---|---|"
    assert lineas[2] == "| a\\|b c | 1.5 |"
    assert lineas[3] == "|  |  |"


# --- Limpieza de lo que contesta el modelo ----------------------------------

@pytest.mark.parametrize("posicion, esperado", [
    ('"bottom"', 'legend=dict(orientation="h", yanchor="bottom", y=-0.25'),
    ("'top center'", 'legend=dict(orientation="h", yanchor="top", y=1.1'),
    ('"right"', 'legend=dict(orientation="v", yanchor="top", y=1'),
])
def test_legend_position_inventado_se_cambia_por_el_dict_de_plotly(posicion, esperado):
    corregido = pipeline.sanitizar_codigo(f"fig.update_layout(legend_position={posicion})")

    assert "legend_position" not in corregido
    assert esperado in corregido


def test_se_extrae_el_bloque_de_codigo_del_lenguaje_pedido():
    assert pipeline.limpiar_bloque_de_codigo(_bloque("x = 1"), "python") == "x = 1"
    assert pipeline.limpiar_bloque_de_codigo(_bloque('{"a": 1}', "json"), "json") == '{"a": 1}'


def test_sin_cercas_de_markdown_se_devuelve_el_texto_tal_cual():
    assert pipeline.limpiar_bloque_de_codigo("  x = 1  ", "python") == "x = 1"


def test_la_respuesta_por_partes_se_une_en_un_solo_texto():
    respuesta = AIMessage(content=[{"type": "text", "text": "hola "}, {"type": "text", "text": "mundo"}, "suelto"])

    assert pipeline.texto_de_respuesta(respuesta) == "hola mundo"


def test_el_contexto_del_fallo_marca_la_linea_de_app_py():
    codigo = "\n".join(f"linea_{n} = {n}" for n in range(1, 31))
    traza = ('File "/usr/lib/python3/site-packages/pandas/core.py", line 999\n'
             'File "/tmp/abc/app.py", line 12, in <module>\n'
             'File "/usr/lib/python3/site-packages/pandas/frame.py", line 5000')

    contexto = pipeline.contexto_del_fallo(codigo, traza)

    assert "--> [LÍNEA DEL FALLO]   12 | linea_12 = 12" in contexto
    assert "     4 | linea_4 = 4" in contexto
    assert "linea_3 = 3" not in contexto


def test_el_contexto_del_fallo_entiende_el_mensaje_de_sintaxis():
    contexto = pipeline.contexto_del_fallo("a = 1\nb = (\nc = 3", "SyntaxError en línea 2, columna 5")

    assert "--> [LÍNEA DEL FALLO]    2 | b = (" in contexto


@pytest.mark.parametrize("mensaje", ["sin número de línea", "line 999"])
def test_sin_una_linea_valida_no_hay_contexto(mensaje):
    assert pipeline.contexto_del_fallo("a = 1", mensaje) == ""


# --- requirements.txt -------------------------------------------------------

def test_requirements_trae_lo_importado_con_la_version_validada():
    codigo = "import os\nimport sklearn.linear_model\nfrom PIL import Image\nimport streamlit as st\nfrom . import local\n"

    requisitos = pipeline.extraer_dependencias(codigo).splitlines()

    assert f"streamlit=={metadata.version('streamlit')}" in requisitos
    assert f"pandas=={metadata.version('pandas')}" in requisitos
    assert "scikit-learn" in [r.split("==")[0] for r in requisitos]
    assert "pillow" in [r.split("==")[0] for r in requisitos]
    assert not any(r.startswith("os") for r in requisitos)
    assert requisitos == sorted(requisitos)


def test_requirements_de_codigo_roto_trae_los_paquetes_base():
    nombres = [r.split("==")[0] for r in pipeline.extraer_dependencias("def (").splitlines()]

    assert nombres == ["jinja2", "matplotlib", "openpyxl", "pandas", "plotly", "streamlit"]


# --- Validación del código generado -----------------------------------------

def test_un_error_de_sintaxis_dice_la_linea():
    ok, mensaje = pipeline.validar_codigo("import streamlit as st\nst.write(", simular=None)

    assert ok is False
    assert "SyntaxError en línea 2" in mensaje


def test_un_byte_nulo_se_reporta_sin_tronar():
    # Python 3.11 (el del Action) lo lanza como ValueError y 3.12+ como
    # SyntaxError; en los dos casos tiene que volver como mensaje, no tronar.
    ok, mensaje = pipeline.validar_codigo("x = 1\x00", simular=None)

    assert ok is False
    assert "null bytes" in mensaje


def test_las_rutas_de_colab_se_rechazan_antes_de_ejecutar():
    ok, mensaje = pipeline.validar_codigo('import pandas as pd\npd.read_excel("/content/drive/MyDrive/a.xlsx")', simular=None)

    assert ok is False
    assert "/content/" in mensaje


def test_apptest_aprueba_una_app_sana():
    assert pipeline.simular_con_apptest(APP_BUENA) == (True, "Simulación AppTest exitosa sin excepciones internas.")


def test_apptest_reporta_la_excepcion_con_su_linea():
    ok, mensaje = pipeline.simular_con_apptest(APP_QUE_TRUENA)

    assert ok is False
    assert "ZeroDivisionError" in mensaje
    assert "--> [LÍNEA DEL FALLO]    3" in pipeline.contexto_del_fallo(APP_QUE_TRUENA, mensaje)


def test_apptest_reporta_un_st_error():
    ok, mensaje = pipeline.simular_con_apptest('import streamlit as st\nst.error("No cuadra el total")\n')

    assert ok is False
    assert "Alerta st.error en UI: No cuadra el total" in mensaje


def test_un_print_de_la_app_no_confunde_el_resultado():
    ok, _ = pipeline.simular_con_apptest('import streamlit as st\nprint("{ruido}")\nst.write("ok")\n')

    assert ok is True


def test_una_app_que_no_termina_se_corta_por_tiempo():
    ok, mensaje = pipeline.simular_con_apptest("import time\ntime.sleep(30)\n", timeout=1)

    assert ok is False
    assert "Fallo de ejecución" in mensaje


def test_el_codigo_generado_no_ve_las_claves(monkeypatch):
    """Un documento con una instrucción escondida no puede sacar los secretos.

    El código que escribe el modelo corre con un entorno mínimo: si viera la
    clave, esta app pintaría un `st.error` y AppTest lo reportaría.
    """
    monkeypatch.setenv("GEMINI_API_KEY_AGENT_1", "clave-secreta")
    monkeypatch.setenv("GITHUB_TOKEN", "token-secreto")
    app = ("import os\nimport streamlit as st\n"
           "fuga = [k for k in os.environ if 'GEMINI' in k or 'GITHUB' in k]\n"
           "if fuga:\n    st.error('FUGA ' + ','.join(fuga))\n")

    assert pipeline.simular_con_apptest(app)[0] is True
    assert "GEMINI_API_KEY_AGENT_1" not in pipeline.entorno_limpio()
    assert "GITHUB_TOKEN" not in pipeline.entorno_limpio()


def test_sin_usuario_aislado_el_comando_corre_tal_cual(monkeypatch):
    monkeypatch.delenv("USUARIO_AISLADO", raising=False)

    assert pipeline.comando_aislado(["python", "app.py"], "/tmp/casa") == ["python", "app.py"]


def test_con_usuario_aislado_el_codigo_corre_como_otro_usuario_y_sin_claves(monkeypatch):
    """En el Action el código generado corre como un usuario sin sudo.

    Con el mismo usuario del pipeline podría leer `/proc/<pid>/environ` del
    proceso padre (las claves) y, en los runners de GitHub, usar `sudo` sin
    contraseña para leer todo el job. Con otro usuario, ninguna de las dos.
    """
    monkeypatch.setenv("USUARIO_AISLADO", "pwo-sandbox")
    monkeypatch.setenv("GEMINI_API_KEY_AGENT_1", "clave-secreta")

    comando = pipeline.comando_aislado(["python", "app.py"], "/tmp/casa")

    assert comando[:7] == ["sudo", "-n", "-u", "pwo-sandbox", "--", "env", "-i"]
    assert "HOME=/tmp/casa" in comando
    assert comando[-2:] == ["python", "app.py"]
    assert not any("GEMINI" in parte or "clave-secreta" in parte for parte in comando)


def test_un_usuario_aislado_que_no_se_puede_usar_detiene_el_pipeline(monkeypatch):
    monkeypatch.setenv("USUARIO_AISLADO", "usuario-que-no-existe-pwo")

    with pytest.raises(pipeline.ConfiguracionIncompleta, match="usuario-que-no-existe-pwo"):
        pipeline.verificar_aislamiento()


def test_sin_sudo_instalado_el_aislamiento_no_arranca(monkeypatch):
    monkeypatch.setenv("USUARIO_AISLADO", "pwo-sandbox")
    monkeypatch.setenv("PATH", "")

    with pytest.raises(pipeline.ConfiguracionIncompleta, match="pwo-sandbox no se puede usar"):
        pipeline.verificar_aislamiento()


def test_el_usuario_aislado_puede_entrar_a_la_carpeta_de_la_app(monkeypatch, tmp_path):
    monkeypatch.setenv("USUARIO_AISLADO", "pwo-sandbox")
    tmp_path.chmod(0o700)

    casa = pipeline.preparar_carpeta(tmp_path)

    assert tmp_path.stat().st_mode & 0o777 == 0o755
    assert (tmp_path / "casa").stat().st_mode & 0o777 == 0o777
    assert casa == str(tmp_path / "casa")


def test_sin_aislar_la_carpeta_conserva_sus_permisos(monkeypatch, tmp_path):
    monkeypatch.delenv("USUARIO_AISLADO", raising=False)
    tmp_path.chmod(0o700)

    pipeline.preparar_carpeta(tmp_path)

    assert tmp_path.stat().st_mode & 0o777 == 0o700


def test_sin_usuario_aislado_no_hay_nada_que_verificar(monkeypatch):
    monkeypatch.delenv("USUARIO_AISLADO", raising=False)

    assert pipeline.verificar_aislamiento() is None


# --- El grafo completo, con el modelo simulado ------------------------------

def _ejecutar(tmp_path: Path, modelos: pipeline.Modelos, navegador=_navegador_aprueba) -> Dict[str, Any]:
    validadores = pipeline.Validadores(simular=pipeline.simular_con_apptest, navegador=navegador)
    return pipeline.ejecutar(str(_documento(tmp_path)), modelos, tmp_path / "salida", validadores)


def test_el_flujo_completo_deja_app_y_requirements_listos(tmp_path):
    desarrollador = ModeloGuionado([_bloque(APP_BUENA)])

    estado = _ejecutar(tmp_path, _modelos(desarrollador))

    assert estado["commit_status"]["status_code"] == pipeline.ESTADO_PUBLICADO
    assert (tmp_path / "salida" / "app.py").read_text(encoding="utf-8") == APP_BUENA
    requisitos = (tmp_path / "salida" / "requirements.txt").read_text(encoding="utf-8")
    assert f"streamlit=={metadata.version('streamlit')}" in requisitos
    assert estado["retry_count"] == 1


def test_los_tres_agentes_reciben_lo_que_les_toca(tmp_path):
    extractor = ModeloGuionado([_bloque(JSON_EXTRAIDO, "json")])
    modelos = _modelos(ModeloGuionado([_bloque(APP_BUENA)]), extractor=extractor)

    _ejecutar(tmp_path, modelos)

    assert "Cimentación" in extractor.recibidos[0][1].content
    assert JSON_EXTRAIDO in modelos.ux.recibidos[0][0].content
    prompt_dev = modelos.desarrollador.recibidos[0][0].content
    assert JSON_EXTRAIDO in prompt_dev
    assert "Pestañas por tabla y métricas arriba." in prompt_dev
    assert 'width="stretch"' in prompt_dev
    assert "use_container_width" not in prompt_dev


def test_un_fallo_vuelve_al_desarrollador_con_el_error_y_el_codigo_anterior(tmp_path):
    desarrollador = ModeloGuionado([_bloque(APP_QUE_TRUENA), _bloque(APP_BUENA)])

    estado = _ejecutar(tmp_path, _modelos(desarrollador))

    segundo_prompt = desarrollador.recibidos[1][0].content
    assert "--- INTENTO FALLIDO 1 ---" in segundo_prompt
    assert "ZeroDivisionError" in segundo_prompt
    assert "--> [LÍNEA DEL FALLO]" in segundo_prompt
    assert "CÓDIGO ANTERIOR QUE GENERÓ EL FALLO" in segundo_prompt
    assert "valor = 1 / 0" in segundo_prompt
    assert estado["retry_count"] == 2
    assert (tmp_path / "salida" / "app.py").read_text(encoding="utf-8") == APP_BUENA


def test_tras_cuatro_intentos_fallidos_la_liga_conserva_el_dashboard_anterior(tmp_path):
    salida = tmp_path / "salida"
    salida.mkdir()
    (salida / "app.py").write_text("# dashboard anterior que funciona\n", encoding="utf-8")
    desarrollador = ModeloGuionado([_bloque(APP_QUE_TRUENA)])

    estado = _ejecutar(tmp_path, _modelos(desarrollador))

    assert len(desarrollador.recibidos) == pipeline.MAX_INTENTOS
    assert estado["commit_status"]["status_code"] == pipeline.ESTADO_ABORTADO
    assert "tras 4 intentos" in estado["commit_status"]["message"]
    assert (salida / "app.py").read_text(encoding="utf-8") == "# dashboard anterior que funciona\n"
    assert not (salida / "requirements.txt").exists()


def test_la_captura_del_navegador_viaja_al_desarrollador(tmp_path):
    respuestas = iter([(False, "Excepción en UI:\nKeyError line 2", "Q0FQVFVSQQ=="), (True, "ok", None)])
    desarrollador = ModeloGuionado([_bloque(APP_BUENA)])

    estado = _ejecutar(tmp_path, _modelos(desarrollador), navegador=lambda codigo: next(respuestas))

    segundo = desarrollador.recibidos[1][0].content
    assert segundo[1]["image_url"]["url"] == "data:image/png;base64,Q0FQVFVSQQ=="
    assert "[Fase 2 - Renderizado Playwright UI]" in segundo[0]["text"]
    assert estado["commit_status"]["status_code"] == pipeline.ESTADO_PUBLICADO


def test_el_legend_position_se_corrige_antes_de_publicar(tmp_path):
    app = 'import streamlit as st\nlayout = dict(legend_position="bottom")\nst.write(layout)\n'

    _ejecutar(tmp_path, _modelos(ModeloGuionado([_bloque(app)])))

    publicado = (tmp_path / "salida" / "app.py").read_text(encoding="utf-8")
    assert "legend_position" not in publicado
    assert 'legend=dict(orientation="h"' in publicado


def test_una_imagen_llega_al_extractor_como_mensaje_multimodal(tmp_path):
    ruta = tmp_path / "foto.png"
    ruta.write_bytes(b"\x89PNG\r\n")
    modelos = _modelos(ModeloGuionado([_bloque(APP_BUENA)]))
    validadores = pipeline.Validadores(simular=pipeline.simular_con_apptest, navegador=_navegador_aprueba)

    pipeline.ejecutar(str(ruta), modelos, tmp_path / "salida", validadores)

    contenido = modelos.extractor.recibidos[0][1].content
    assert contenido[1]["image_url"]["url"].startswith("data:image/png;base64,")


# --- Configuración de los modelos -------------------------------------------

def test_sin_claves_de_gemini_se_dice_cuales_faltan():
    with pytest.raises(pipeline.ConfiguracionIncompleta, match="GEMINI_API_KEY_AGENT_2"):
        pipeline.construir_modelos({"GEMINI_API_KEY_AGENT_1": "a", "GEMINI_API_KEY_AGENT_3": "c"})


def test_una_sola_clave_sirve_para_los_tres_agentes_y_el_modelo_se_cambia_por_entorno():
    modelos = pipeline.construir_modelos({"GEMINI_API_KEY": "k", "GEMINI_MODEL": "gemini-x-prueba"})

    for llm in (modelos.extractor, modelos.ux, modelos.desarrollador):
        assert llm.model.endswith("gemini-x-prueba")


def test_sin_gemini_model_se_usa_el_modelo_de_gs3():
    modelos = pipeline.construir_modelos({"GEMINI_API_KEY": "k"})

    assert modelos.extractor.model.endswith(pipeline.MODELO_POR_DEFECTO)


# --- Descarga del documento desde Supabase Storage --------------------------

class _Storage(http.server.BaseHTTPRequestHandler):
    """Un Storage de mentiras en 127.0.0.1: sirve archivos, 404 y redirecciones."""

    archivos: Dict[str, bytes] = {}

    def do_GET(self) -> None:
        if self.path.endswith("/redirige.xlsx"):
            self.send_response(302)
            self.send_header("Location", "http://ejemplo.invalid/otro.xlsx")
            self.end_headers()
            return
        cuerpo = self.archivos.get(self.path)
        self.send_response(200 if cuerpo is not None else 404)
        self.end_headers()
        self.wfile.write(cuerpo or b"")

    def log_message(self, *args):
        return


@pytest.fixture
def storage():
    servidor = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Storage)
    hilo = threading.Thread(target=servidor.serve_forever, daemon=True)
    hilo.start()
    base = f"http://127.0.0.1:{servidor.server_address[1]}"
    _Storage.archivos = {"/storage/v1/object/public/archivos/2026/Junta%20HLT.xlsx": b"EXCEL"}
    yield base, f"{base}/storage/v1/object/public/"
    servidor.shutdown()


def test_descarga_el_documento_permitido_con_un_nombre_seguro(storage, tmp_path):
    base, origen = storage

    ruta = pipeline.descargar_archivo(f"{base}/storage/v1/object/public/archivos/2026/Junta%20HLT.xlsx", tmp_path, origen)

    assert ruta == tmp_path / "Junta_HLT.xlsx"
    assert ruta.read_bytes() == b"EXCEL"


@pytest.mark.parametrize("url", [
    "http://127.0.0.2:1/storage/v1/object/public/archivos/a.xlsx",
    "{base}/storage/v1/object/sign/archivos/a.xlsx",
    "{base}/storage/v1/object/public/../sign/archivos/a.xlsx",
    "https://{host}/storage/v1/object/public/archivos/a.xlsx",
])
def test_no_se_descarga_nada_fuera_del_storage_de_holtmont(storage, tmp_path, url):
    base, origen = storage

    with pytest.raises(pipeline.OrigenNoPermitido):
        pipeline.descargar_archivo(url.format(base=base, host=base.split("//")[1]), tmp_path, origen)


def test_sin_origen_configurado_no_se_descarga_nada(tmp_path):
    with pytest.raises(pipeline.ConfiguracionIncompleta, match="ORIGEN_PERMITIDO"):
        pipeline.descargar_archivo("https://x.supabase.co/storage/v1/object/public/a.xlsx", tmp_path, "")


def test_un_formato_que_el_pipeline_no_lee_se_rechaza(storage, tmp_path):
    base, origen = storage

    with pytest.raises(pipeline.ArchivoIlegible, match=r"\.exe"):
        pipeline.descargar_archivo(f"{base}/storage/v1/object/public/archivos/virus.exe", tmp_path, origen)


def test_un_documento_demasiado_grande_se_rechaza(storage, tmp_path):
    base, origen = storage

    with pytest.raises(pipeline.ArchivoIlegible, match="pesa más de 3 bytes"):
        pipeline.descargar_archivo(f"{base}/storage/v1/object/public/archivos/2026/Junta%20HLT.xlsx",
                                   tmp_path, origen, max_bytes=3)


def test_una_redireccion_no_se_sigue(storage, tmp_path):
    base, origen = storage

    with pytest.raises(pipeline.OrigenNoPermitido, match="redirige"):
        pipeline.descargar_archivo(f"{base}/storage/v1/object/public/archivos/redirige.xlsx", tmp_path, origen)


def test_un_documento_que_no_existe_se_reporta(storage, tmp_path):
    base, origen = storage

    with pytest.raises(pipeline.ArchivoIlegible, match="404"):
        pipeline.descargar_archivo(f"{base}/storage/v1/object/public/archivos/falta.xlsx", tmp_path, origen)


def test_un_storage_caido_se_reporta(tmp_path):
    with pytest.raises(pipeline.ArchivoIlegible, match="No se pudo descargar"):
        pipeline.descargar_archivo("http://127.0.0.1:9/storage/v1/object/public/a.xlsx", tmp_path,
                                   "http://127.0.0.1:9/storage/v1/object/public/")


# --- La línea de comandos que corre el Action -------------------------------

def test_main_publica_y_sale_en_cero(tmp_path, capsys):
    modelos = _modelos(ModeloGuionado([_bloque(APP_BUENA)]))
    validadores = pipeline.Validadores(simular=pipeline.simular_con_apptest, navegador=_navegador_aprueba)

    codigo = pipeline.main(["--archivo", str(_documento(tmp_path)), "--salida", str(tmp_path / "salida"),
                            "--folio", "1234XX Const 061026"], construir=lambda: modelos, validadores=validadores)

    assert codigo == 0
    assert (tmp_path / "salida" / "app.py").exists()
    salida = capsys.readouterr().out
    assert "Folio: 1234XX Const 061026" in salida
    assert pipeline.ESTADO_PUBLICADO in salida


def test_main_sale_en_uno_si_se_aborta(tmp_path):
    modelos = _modelos(ModeloGuionado([_bloque(APP_QUE_TRUENA)]))
    validadores = pipeline.Validadores(simular=pipeline.simular_con_apptest, navegador=_navegador_aprueba)

    codigo = pipeline.main(["--archivo", str(_documento(tmp_path)), "--salida", str(tmp_path / "salida")],
                           construir=lambda: modelos, validadores=validadores)

    assert codigo == 1
    assert not (tmp_path / "salida" / "app.py").exists()


def test_main_descarga_por_url_con_el_origen_del_entorno(storage, tmp_path, monkeypatch):
    base, origen = storage
    _Storage.archivos["/storage/v1/object/public/archivos/datos.csv"] = "concepto,monto\nMuro,10\n".encode()
    monkeypatch.setenv("ORIGEN_PERMITIDO", origen)
    extractor = ModeloGuionado([_bloque(JSON_EXTRAIDO, "json")])
    modelos = _modelos(ModeloGuionado([_bloque(APP_BUENA)]), extractor=extractor)
    validadores = pipeline.Validadores(simular=pipeline.simular_con_apptest, navegador=_navegador_aprueba)

    codigo = pipeline.main(["--url", f"{base}/storage/v1/object/public/archivos/datos.csv",
                            "--salida", str(tmp_path / "salida")], construir=lambda: modelos, validadores=validadores)

    assert codigo == 0
    assert "| Muro | 10 |" in extractor.recibidos[0][1].content


def test_main_no_arranca_si_el_aislamiento_no_sirve(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("USUARIO_AISLADO", "usuario-que-no-existe-pwo")

    def no_debe_llamarse():
        raise AssertionError("sin aislamiento no se gasta ni una llamada al modelo")

    codigo = pipeline.main(["--archivo", str(_documento(tmp_path)), "--salida", str(tmp_path / "salida")],
                           construir=no_debe_llamarse)

    assert codigo == 1
    assert "USUARIO_AISLADO" in capsys.readouterr().err


def test_main_sin_claves_falla_con_mensaje_claro(tmp_path, monkeypatch, capsys):
    for variable in ("GEMINI_API_KEY", "GEMINI_API_KEY_AGENT_1", "GEMINI_API_KEY_AGENT_2", "GEMINI_API_KEY_AGENT_3"):
        monkeypatch.delenv(variable, raising=False)

    codigo = pipeline.main(["--archivo", str(_documento(tmp_path)), "--salida", str(tmp_path / "salida")])

    assert codigo == 1
    assert "Faltan las claves de Gemini" in capsys.readouterr().err
