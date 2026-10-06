# -*- coding: utf-8 -*-
"""
Dashboard de la Pre Work Order: de un documento a una app de Streamlit publicada.

Adaptación de `gs3.py` (GS3_POTENCIADO, Colab) para que corra sola en el GitHub
Action del repositorio `Agente_Full` cada vez que alguien sube un documento a la
Pre Work Order de Holtmont:

    PWO (index.html) --sube xlsx/pdf/docx--> Supabase Storage
      -> api/services/dashboard_pwo.py manda un `repository_dispatch`
      -> el Action de Agente_Full corre ESTE archivo:
           Agente 1 extrae JSON -> Agente 2 diseña -> Agente 3 programa
           -> sintaxis + AppTest + Playwright (hasta 4 intentos)
           -> escribe app.py + requirements.txt
      -> el Action hace el commit y Streamlit Cloud vuelve a desplegar
      -> la liga corta (subdominio fijo de Streamlit) enseña el dashboard nuevo

Qué cambió respecto de gs3.py, y por qué:

* **Credenciales fuera del código.** gs3.py traía un token de GitHub y tres API
  keys de Gemini en claro. Aquí se leen del entorno (secretos del Action) y el
  commit lo hace el `GITHUB_TOKEN` del propio Action, no un token personal.
* **El código que escribe el modelo corre sin secretos.** AppTest se ejecutaba
  en el mismo proceso que tenía las claves en `os.environ`: un documento con una
  instrucción escondida podía hacer que el dashboard público las enseñara.
  Ahora AppTest y `streamlit run` van en un subproceso con `entorno_limpio()`
  y, en el Action, como un usuario sin sudo (`USUARIO_AISLADO`).
* **El documento llega por URL y solo del Storage de Holtmont**
  (`ORIGEN_PERMITIDO`), con tope de tamaño y sin seguir redirecciones.
* **Un documento ilegible detiene el pipeline.** gs3.py le pasaba al modelo el
  texto "Error procesando archivo: ..." y publicaba un dashboard sobre el error.
* **Si se agotan los intentos, `app.py` no se toca:** la liga sigue enseñando el
  último dashboard bueno y el Action termina en rojo.
* **`requirements.txt` fija las versiones con las que se validó la app**, para
  que Streamlit Cloud instale lo mismo que se probó.
* **Los selectores de error en el navegador son los de Streamlit 1.65**
  (`stAlertContentError`); los de gs3.py (`stAlertErrorIcon`, `.alert-danger`)
  ya no existen y un `st.error` pasaba la auditoría sin que nadie lo viera.

Uso (lo mismo en el Action que a mano):

    python pipeline_dashboard.py --url "<url pública del archivo>" --salida .
    python pipeline_dashboard.py --archivo junta.xlsx --salida /tmp/dashboard
"""

from __future__ import annotations

import argparse
import ast
import base64
import concurrent.futures
import functools
import json
import os
import posixpath
import re
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from importlib import metadata
from pathlib import Path
from typing import Any, Callable, Dict, List, Mapping, Optional, Tuple, TypedDict, Union

import pandas as pd
from langchain_core.messages import HumanMessage, SystemMessage
from langgraph.graph import END, START, StateGraph

MAX_INTENTOS = 4
MAX_BYTES_ARCHIVO = 20 * 1024 * 1024
# El que usaba gs3.py. `GEMINI_MODEL` lo cambia sin tocar código: los
# proveedores retiran modelos (ver api/modelos_llm.py en HOLTMONT-PYTHON).
MODELO_POR_DEFECTO = "gemini-3.5-flash"
EXTENSIONES_SOPORTADAS = ("csv", "xls", "xlsx", "pdf", "docx", "png", "jpg", "jpeg", "txt")

ESTADO_PUBLICADO = "LISTO_PARA_COMMIT"
ESTADO_ABORTADO = "ABORTED_VALIDATION_ERROR"

# Lo único del entorno que ve el código generado. Lista blanca y no negra: un
# secreto nuevo que alguien añada al Action queda fuera sin tocar este archivo.
VARIABLES_DEL_SUBPROCESO = ("PATH", "HOME", "LANG", "LC_ALL", "TMPDIR")

Contenido = Union[str, List[Dict[str, Any]]]


class ErrorDelPipeline(Exception):
    """Fallo que termina el pipeline con un mensaje para la bitácora del Action."""


class ConfiguracionIncompleta(ErrorDelPipeline):
    """Falta una variable de entorno sin la cual no se puede trabajar."""


class OrigenNoPermitido(ErrorDelPipeline):
    """La URL del documento no apunta al Storage de Holtmont."""


class ArchivoIlegible(ErrorDelPipeline):
    """El documento no se pudo descargar o leer."""


class GraphState(TypedDict, total=False):
    file_path: str
    extracted_json: str
    ux_prompt: str
    python_code: str
    requirements_txt: str
    validation_error: Optional[str]
    error_history: List[str]
    ui_screenshot: Optional[str]
    retry_count: int
    is_valid: bool
    commit_status: Dict[str, Any]


# --- Lectura del documento --------------------------------------------------

INSTRUCCION_IMAGEN = (
    "Analiza esta imagen minuciosamente. Extrae todos los datos visibles, "
    "incluyendo encabezados, logotipos, fechas, tablas completas, totales, "
    "subtotales, métricas aisladas y notas con su significado contextual."
)


def _celda(valor: Any) -> str:
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    return str(valor).replace("|", "\\|").replace("\r", " ").replace("\n", " ").strip()


def tabla_markdown(df: pd.DataFrame) -> str:
    """El DataFrame como tabla markdown, sin depender de `tabulate`."""
    encabezados = [_celda(c) for c in df.columns]
    lineas = ["| " + " | ".join(encabezados) + " |", "|" + "---|" * len(encabezados)]
    for fila in df.itertuples(index=False, name=None):
        lineas.append("| " + " | ".join(_celda(v) for v in fila) + " |")
    return "\n".join(lineas)


def _sin_vacios(df: pd.DataFrame) -> pd.DataFrame:
    return df.dropna(how="all").dropna(axis=1, how="all")


def _leer_csv(ruta: str) -> str:
    try:
        df = pd.read_csv(ruta, encoding="utf-8-sig", nrows=1000)
    except UnicodeDecodeError:
        df = pd.read_csv(ruta, encoding="latin-1", nrows=1000)
    return tabla_markdown(_sin_vacios(df))


def _leer_excel(ruta: str) -> str:
    libro = pd.ExcelFile(ruta)
    secciones = []
    for hoja in libro.sheet_names:
        df = pd.read_excel(libro, sheet_name=hoja, nrows=500)
        secciones.append(f"### PESTAÑA/HOJA EXCEL: {hoja}\n{tabla_markdown(_sin_vacios(df))}")
    return "\n\n".join(secciones)


def _leer_pdf(ruta: str) -> str:
    from pypdf import PdfReader

    paginas = PdfReader(ruta).pages[:20]
    return "\n".join(f"--- PÁGINA {i + 1} ---\n{pagina.extract_text() or ''}" for i, pagina in enumerate(paginas))


def _leer_docx(ruta: str) -> str:
    import docx

    documento = docx.Document(ruta)
    elementos = [p.text.strip() for p in documento.paragraphs if p.text.strip()]
    for i, tabla in enumerate(documento.tables):
        elementos.append(f"\n--- TABLA {i + 1} DE WORD ---")
        for fila in tabla.rows:
            elementos.append(" | ".join(c.text.strip().replace("\n", " ") for c in fila.cells))
    return "\n".join(elementos)


def _leer_doc(ruta: str) -> str:
    raise ValueError("el formato .doc no se puede leer; guárdalo como .docx")


def _leer_imagen(ruta: str) -> List[Dict[str, Any]]:
    extension = Path(ruta).suffix.lower().lstrip(".")
    mime = "jpeg" if extension in ("jpg", "jpeg") else extension
    codificada = base64.b64encode(Path(ruta).read_bytes()).decode("ascii")
    return [
        {"type": "text", "text": INSTRUCCION_IMAGEN},
        {"type": "image_url", "image_url": {"url": f"data:image/{mime};base64,{codificada}"}},
    ]


def _leer_texto(ruta: str) -> str:
    with open(ruta, "r", encoding="utf-8", errors="ignore") as archivo:
        return archivo.read()[:10000]


_LECTORES: Dict[str, Callable[[str], Contenido]] = {
    "csv": _leer_csv, "xls": _leer_excel, "xlsx": _leer_excel, "pdf": _leer_pdf,
    "docx": _leer_docx, "doc": _leer_doc, "png": _leer_imagen, "jpg": _leer_imagen,
    "jpeg": _leer_imagen,
}


def leer_archivo(ruta: str) -> Contenido:
    """Texto, tablas o imagen del documento, listo para el Agente 1."""
    extension = Path(ruta).suffix.lower().lstrip(".")
    try:
        return _LECTORES.get(extension, _leer_texto)(ruta)
    except Exception as exc:
        raise ArchivoIlegible(f"No se pudo leer {Path(ruta).name}: {exc}") from exc


# --- Lo que contesta el modelo ----------------------------------------------

_LEYENDAS = (
    (r"bottom|bottom\s+center|below", 'legend=dict(orientation="h", yanchor="bottom", y=-0.25, xanchor="center", x=0.5)'),
    (r"top|top\s+center|above", 'legend=dict(orientation="h", yanchor="top", y=1.1, xanchor="center", x=0.5)'),
    (r"right", 'legend=dict(orientation="v", yanchor="top", y=1, xanchor="left", x=1.02)'),
)


def sanitizar_codigo(codigo: str) -> str:
    """Corrige el `legend_position` que el modelo inventa para Plotly."""
    for posiciones, reemplazo in _LEYENDAS:
        patron = rf'legend_position\s*=\s*["\']({posiciones})["\']'
        codigo = re.sub(patron, reemplazo, codigo, flags=re.IGNORECASE)
    return codigo


def limpiar_bloque_de_codigo(texto: str, lenguaje: str = "python") -> str:
    """El contenido del bloque ```lenguaje```, sin la prosa que lo rodea."""
    coincidencia = re.search(rf"```(?:{lenguaje})?[ \t]*\n?(.*?)```", texto, re.DOTALL | re.IGNORECASE)
    if coincidencia:
        return coincidencia.group(1).strip()
    return texto.replace(f"```{lenguaje}", "").replace("```", "").strip()


def texto_de_respuesta(respuesta: Any) -> str:
    contenido = respuesta.content
    if isinstance(contenido, list):
        return "".join(parte.get("text", "") for parte in contenido if isinstance(parte, dict))
    return str(contenido)


def _linea_del_fallo(mensaje: str) -> Optional[int]:
    """La línea de `app.py` donde truena; si la traza no la nombra, la última citada."""
    propias = re.findall(r'app\.py", line (\d+)', mensaje)
    citadas = propias or re.findall(r"(?:line\s+|línea\s+)(\d+)", mensaje, re.IGNORECASE)
    return int(citadas[-1]) if citadas else None


def contexto_del_fallo(codigo: str, mensaje: str, ventana: int = 8) -> str:
    """El fragmento de código alrededor de la línea que falló, marcada."""
    numero = _linea_del_fallo(mensaje)
    lineas = codigo.splitlines()
    if numero is None or not 1 <= numero <= len(lineas):
        return ""
    fragmento = []
    for indice in range(max(1, numero - ventana), min(len(lineas), numero + ventana) + 1):
        marca = "--> [LÍNEA DEL FALLO] " if indice == numero else " " * 22
        fragmento.append(f"{marca}{indice:4d} | {lineas[indice - 1]}")
    return f"\n\nTRAZABILIDAD DIRECTA DEL CÓDIGO (LÍNEA {numero}):\n" + "\n".join(fragmento)


# --- requirements.txt -------------------------------------------------------

_PAQUETES_BASE = ("streamlit", "pandas", "plotly", "openpyxl", "jinja2", "matplotlib")
_NOMBRE_DEL_PAQUETE = {
    "PIL": "pillow", "sklearn": "scikit-learn", "cv2": "opencv-python", "yaml": "pyyaml",
    "docx": "python-docx", "bs4": "beautifulsoup4",
}


def _requisito(paquete: str) -> str:
    """`paquete==versión` con la que se validó; sin versión si no está instalado."""
    try:
        return f"{paquete}=={metadata.version(paquete)}"
    except metadata.PackageNotFoundError:
        return paquete


def _modulos_importados(arbol: ast.AST) -> set:
    modulos = set()
    for nodo in ast.walk(arbol):
        if isinstance(nodo, ast.Import):
            modulos.update(alias.name.split(".")[0] for alias in nodo.names)
        elif isinstance(nodo, ast.ImportFrom) and nodo.module and not nodo.level:
            modulos.add(nodo.module.split(".")[0])
    return modulos


def extraer_dependencias(codigo: str) -> str:
    """El requirements.txt de la app, con las versiones del entorno que la validó."""
    paquetes = set(_PAQUETES_BASE)
    try:
        arbol = ast.parse(codigo)
    except SyntaxError:
        arbol = None
    if arbol is not None:
        estandar = getattr(sys, "stdlib_module_names", set())
        paquetes.update(_NOMBRE_DEL_PAQUETE.get(m, m) for m in _modulos_importados(arbol) if m not in estandar)
    return "\n".join(sorted(_requisito(p) for p in paquetes)) + "\n"


# --- Validación: sintaxis, rutas y AppTest ----------------------------------

def entorno_limpio(casa: Optional[str] = None) -> Dict[str, str]:
    """El entorno con el que corre el código generado: sin claves ni tokens.

    `casa` sustituye a HOME solo para el usuario aislado, que no tiene una
    propia. Sin aislar se conserva la HOME real: cambiarla esconde los paquetes
    instalados con `pip --user` y la app falla por un import que sí existe.
    """
    entorno = {k: v for k, v in os.environ.items() if k in VARIABLES_DEL_SUBPROCESO}
    entorno["STREAMLIT_BROWSER_GATHER_USAGE_STATS"] = "false"
    if casa:
        entorno["HOME"] = casa
    return entorno


def _usuario_aislado() -> str:
    return os.environ.get("USUARIO_AISLADO", "").strip()


def comando_aislado(comando: List[str], casa: str) -> List[str]:
    """`comando` tal cual o, si el Action define `USUARIO_AISLADO`, como ese usuario.

    Quitar las claves del entorno no basta: con el mismo usuario del pipeline,
    el código generado puede leer `/proc/<pid>/environ` del proceso padre y, en
    los runners de GitHub, usar `sudo` sin contraseña para leer todo el job.
    Otro usuario, sin sudo, no puede ninguna de las dos cosas.
    """
    usuario = _usuario_aislado()
    if not usuario:
        return comando
    variables = [f"{clave}={valor}" for clave, valor in entorno_limpio(casa).items()]
    return ["sudo", "-n", "-u", usuario, "--", "env", "-i", *variables, *comando]


def verificar_aislamiento() -> None:
    """Falla antes de gastar una llamada al modelo si el usuario aislado no sirve."""
    usuario = _usuario_aislado()
    if not usuario:
        return None
    try:
        prueba = subprocess.run(["sudo", "-n", "-u", usuario, "--", "true"], capture_output=True, timeout=30)
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise ConfiguracionIncompleta(f"USUARIO_AISLADO={usuario} no se puede usar: {exc}") from exc
    if prueba.returncode != 0:
        detalle = prueba.stderr.decode("utf-8", errors="ignore").strip()
        raise ConfiguracionIncompleta(f"USUARIO_AISLADO={usuario} no se puede usar: {detalle}")
    return None


def preparar_carpeta(carpeta: Path) -> str:
    """La carpeta `casa` (HOME del código generado), legible por el usuario aislado."""
    casa = carpeta / "casa"
    casa.mkdir(exist_ok=True)
    if _usuario_aislado():
        carpeta.chmod(0o755)
        casa.chmod(0o777)
    return str(casa)


def validar_sintaxis(codigo: str) -> Tuple[bool, str]:
    try:
        compile(codigo, "<app_validation>", "exec")
    except SyntaxError as exc:
        return False, (f"SyntaxError en línea {exc.lineno}, columna {exc.offset}: {exc.msg}\n"
                       f"Fragmento problemático: {exc.text}")
    except (ValueError, TypeError) as exc:
        return False, f"Fallo general de compilación: {exc}"
    return True, "Sintaxis válida."


def validar_rutas(codigo: str) -> Tuple[bool, str]:
    if "/content/" in codigo or "drive/MyDrive" in codigo:
        return False, (
            "Violación de entorno: El script contiene rutas locales fijas ('/content/...'). "
            "El dashboard debe embeber los datos sintéticos como DataFrames directos "
            "o proveer un 'st.sidebar.file_uploader' con fallback automático."
        )
    return True, ""


_MARCA_APPTEST = "__APPTEST__"
_SCRIPT_APPTEST = r'''
import json, sys
from streamlit.testing.v1 import AppTest

at = AppTest.from_file(sys.argv[1], default_timeout=float(sys.argv[2]))
at.run()
excepciones = []
for exc in at.exception:
    # El nombre de la excepción vive en el proto; `exc.type` dice "exception".
    tipo = str(getattr(getattr(exc, "proto", None), "type", "") or "")
    detalle = "Excepción en UI: " + (tipo + ": " if tipo else "") + str(getattr(exc, "value", exc))
    mensaje = str(getattr(exc, "message", "") or "")
    if mensaje and mensaje not in detalle:
        detalle += " (" + mensaje + ")"
    detalle += "".join("\n" + str(t) for t in (getattr(exc, "stack_trace", None) or []))
    excepciones.append(detalle)
alertas = ["Alerta st.error en UI: " + str(getattr(e, "value", e)) for e in at.error]
print("__APPTEST__" + json.dumps({"excepciones": excepciones, "alertas": alertas}))
'''


def _resultado_marcado(salida: str) -> Optional[Dict[str, List[str]]]:
    for linea in reversed(salida.splitlines()):
        if linea.startswith(_MARCA_APPTEST):
            return json.loads(linea[len(_MARCA_APPTEST):])
    return None


def _interpretar_apptest(proceso: subprocess.CompletedProcess) -> Tuple[bool, str]:
    resultado = _resultado_marcado(proceso.stdout or "")
    if resultado is None:
        return False, (f"Fallo de ejecución en tiempo real (código {proceso.returncode}):\n"
                       f"{(proceso.stderr or '')[-4000:]}")
    if resultado["excepciones"]:
        return False, ("Fallo visual detectado en la interfaz de Streamlit (st.exception):\n"
                       + "\n".join(resultado["excepciones"]))
    if resultado["alertas"]:
        return False, "Alertas críticas detectadas en la UI de Streamlit:\n" + "\n".join(resultado["alertas"])
    return True, "Simulación AppTest exitosa sin excepciones internas."


def simular_con_apptest(codigo: str, timeout: float = 15.0) -> Tuple[bool, str]:
    """Corre la app con AppTest en un subproceso sin secretos."""
    with tempfile.TemporaryDirectory() as carpeta:
        ruta = Path(carpeta) / "app.py"
        ruta.write_text(codigo, encoding="utf-8")
        casa = preparar_carpeta(Path(carpeta))
        comando = comando_aislado([sys.executable, "-c", _SCRIPT_APPTEST, str(ruta), str(timeout)], casa)
        try:
            proceso = subprocess.run(comando, capture_output=True, text=True, timeout=timeout + 60,
                                     env=entorno_limpio(), cwd=carpeta)
        except subprocess.TimeoutExpired:
            return False, f"Fallo de ejecución: AppTest no terminó en {timeout + 60:.0f} s."
    return _interpretar_apptest(proceso)


def validar_codigo(codigo: str, simular: Optional[Callable[[str], Tuple[bool, str]]] = simular_con_apptest
                   ) -> Tuple[bool, str]:
    """Fase 1: sintaxis, rutas de Colab y ejecución con AppTest, en ese orden."""
    for paso in (validar_sintaxis, validar_rutas):
        ok, mensaje = paso(codigo)
        if not ok:
            return False, mensaje
    ok, mensaje = simular(codigo)
    if not ok:
        return False, mensaje
    return True, "Sintaxis, dependencias y ejecución en runtime 100% funcionales."


# --- Validación: la app servida y vista en un navegador ---------------------

_SIN_PROXY = urllib.request.build_opener(urllib.request.ProxyHandler({}))


def puerto_libre() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as conexion:
        conexion.bind(("127.0.0.1", 0))
        return conexion.getsockname()[1]


def arrancar_streamlit(ruta: Path, puerto: int, bitacora: Any) -> subprocess.Popen:
    comando = [
        sys.executable, "-m", "streamlit", "run", str(ruta), f"--server.port={puerto}",
        "--server.headless=true", "--browser.gatherUsageStats=false", "--server.enableCORS=false",
        "--server.enableXsrfProtection=false", "--client.showErrorDetails=full",
    ]
    casa = preparar_carpeta(ruta.parent)
    return subprocess.Popen(comando_aislado(comando, casa), stdout=subprocess.DEVNULL, stderr=bitacora,
                            env=entorno_limpio(), cwd=str(ruta.parent))


def _responde(url: str) -> bool:
    try:
        with _SIN_PROXY.open(url, timeout=1) as respuesta:
            return respuesta.status == 200
    except (urllib.error.URLError, OSError):
        return False


def esperar_servidor(proceso: subprocess.Popen, puerto: int, intentos: int = 40,
                     pausa: float = 0.5) -> Tuple[bool, str]:
    salud = f"http://localhost:{puerto}/_stcore/health"
    for _ in range(intentos):
        if proceso.poll() is not None:
            return False, "El servidor Streamlit falló al iniciar en subproceso."
        if _responde(salud):
            return True, ""
        time.sleep(pausa)
    return False, f"Timeout: El servidor de Streamlit no respondió tras {intentos * pausa:.0f} segundos."


def detener(proceso: subprocess.Popen) -> None:
    if proceso.poll() is not None:
        return
    proceso.terminate()
    try:
        proceso.wait(timeout=3)
    except subprocess.TimeoutExpired:
        proceso.kill()
        proceso.wait()


def esperar_inactividad(pagina: Any, timeout_ms: int = 10000) -> None:
    pagina.wait_for_timeout(300)
    try:
        pagina.wait_for_selector('[data-testid="stStatusWidget"]', state="detached", timeout=timeout_ms)
    except Exception:
        pass  # Sin indicador de "corriendo" no hay nada que esperar.
    pagina.wait_for_timeout(300)


def _anadir(errores: List[str], error: str) -> None:
    if error.strip() and error not in errores:
        errores.append(error)


def escanear_errores(pagina: Any) -> List[str]:
    """Excepciones, `st.error` y gráficas de Plotly vacías que se ven en la página."""
    errores: List[str] = []
    for elemento in pagina.query_selector_all('[data-testid="stException"]'):
        _anadir(errores, f"Excepción en UI:\n{elemento.inner_text().strip()}")
    for elemento in pagina.query_selector_all('[data-testid="stAlertContentError"]'):
        _anadir(errores, f"Alerta st.error en UI:\n{elemento.inner_text().strip()}")
    vacias = pagina.query_selector_all(".js-plotly-plot:empty")
    if vacias:
        _anadir(errores, f"Se detectaron {len(vacias)} contenedor(es) de Plotly vacíos en la interfaz.")
    return errores


def _captura(pagina: Any) -> Optional[str]:
    try:
        return base64.b64encode(pagina.screenshot(full_page=True)).decode("ascii")
    except Exception:
        return None


class Auditor:
    """Junta los errores de la página sin repetirlos, con una captura del primero."""

    def __init__(self, pagina: Any) -> None:
        self._pagina = pagina
        self._vistos: set = set()
        self.errores: List[str] = []
        self.captura: Optional[str] = None

    def _anotar(self, contexto: str, error: str) -> None:
        if error in self._vistos:
            return
        self._vistos.add(error)
        self.errores.append(f"[{contexto}]:\n{error}")
        if self.captura is None:
            self.captura = _captura(self._pagina)

    def registrar(self, contexto: str) -> None:
        for error in escanear_errores(self._pagina):
            self._anotar(contexto, error)

    def agregar_consola(self, mensajes: List[str]) -> None:
        criticos = [m for m in mensajes if "favicon" not in m.lower() and "source-map" not in m.lower()]
        for mensaje in criticos[:2]:
            self._anotar("Fallo crítico en consola del navegador", mensaje)

    def veredicto(self) -> Tuple[bool, str, Optional[str]]:
        if not self.errores:
            return True, "Renderizado completo en navegador sin errores visuales tras interactuar con todos los componentes.", None
        return False, (f"Fallo(s) reactivo(s) detectado(s) durante la validación integral de la interfaz "
                       f"({len(self.errores)} inconsistencia(s)):\n" + "\n---\n".join(self.errores)), self.captura


def _abrir_segunda_pestana(pagina: Any) -> bool:
    pestanas = pagina.query_selector_all('[data-testid="stTab"]')
    if len(pestanas) < 2:
        return False
    pestanas[1].click(timeout=2000)
    return True


def _pulsar_boton(pagina: Any) -> bool:
    boton = pagina.query_selector('[data-testid="stButton"] button')
    if not (boton and boton.is_visible()):
        return False
    boton.click(timeout=2000)
    return True


def _mover_slider(pagina: Any) -> bool:
    control = pagina.query_selector('[role="slider"]')
    if not (control and control.is_visible()):
        return False
    control.focus()
    pagina.keyboard.press("ArrowRight")
    pagina.keyboard.press("ArrowRight")
    return True


def _cambiar_radio(pagina: Any) -> bool:
    opciones = pagina.query_selector_all('[data-testid="stRadio"] label')
    if len(opciones) < 2:
        return False
    opciones[1].click(timeout=2000)
    return True


def _elegir_en_lista(pagina: Any, selector: str, preferida: int) -> bool:
    caja = pagina.query_selector(selector)
    if not (caja and caja.is_visible()):
        return False
    caja.click(timeout=2000)
    pagina.wait_for_timeout(300)
    opciones = pagina.query_selector_all('li[role="option"], div[role="option"]')
    if opciones:
        opciones[min(preferida, len(opciones) - 1)].click(timeout=2000)
    pagina.keyboard.press("Escape")
    return True


INTERACCIONES: Tuple[Tuple[str, Callable[[Any], bool]], ...] = (
    ("el renderizado de la pestaña secundaria", _abrir_segunda_pestana),
    ("el clic en botón interactivo", _pulsar_boton),
    ("la interacción con control deslizante", _mover_slider),
    ("el cambio de opción en radio button", _cambiar_radio),
    ("la selección en menú desplegable (selectbox)",
     functools.partial(_elegir_en_lista, selector='[data-testid="stSelectbox"]', preferida=1)),
    ("la selección en selector múltiple (multiselect)",
     functools.partial(_elegir_en_lista, selector='[data-testid="stMultiSelect"]', preferida=0)),
)


def recorrer_controles(pagina: Any, auditor: Auditor) -> None:
    for contexto, accion in INTERACCIONES:
        try:
            if not accion(pagina):
                continue
            esperar_inactividad(pagina)
            auditor.registrar(contexto)
        except Exception:
            # Igual que en gs3.py: un control que no se deja pulsar (tapado, en
            # animación) no es un fallo del dashboard. Lo que se audita es lo
            # que la app pinta después de cada interacción, no el clic.
            continue


def _escuchar_consola(pagina: Any) -> List[str]:
    mensajes: List[str] = []
    pagina.on("pageerror", lambda error: mensajes.append(f"JS Page Error: {error}"))
    pagina.on("console", lambda m: mensajes.append(f"Console Error: {m.text}") if m.type == "error" else None)
    return mensajes


def auditar_pagina(url: str) -> Tuple[bool, str, Optional[str]]:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        navegador = playwright.chromium.launch(headless=True)
        try:
            pagina = navegador.new_context(viewport={"width": 1280, "height": 800}).new_page()
            consola = _escuchar_consola(pagina)
            pagina.goto(url, wait_until="domcontentloaded", timeout=20000)
            pagina.wait_for_selector('[data-testid="stAppViewContainer"]', timeout=15000)
            esperar_inactividad(pagina)
            auditor = Auditor(pagina)
            auditor.registrar("la carga inicial del dashboard")
            recorrer_controles(pagina, auditor)
            auditor.agregar_consola(consola)
            return auditor.veredicto()
        finally:
            navegador.close()


def _servir_y_auditar(proceso: subprocess.Popen, puerto: int) -> Tuple[bool, str, Optional[str]]:
    listo, detalle = esperar_servidor(proceso, puerto)
    if not listo:
        return False, detalle, None
    return auditar_pagina(f"http://localhost:{puerto}")


def _validar_en_navegador(codigo: str) -> Tuple[bool, str, Optional[str]]:
    with tempfile.TemporaryDirectory() as carpeta:
        ruta = Path(carpeta) / "app.py"
        ruta.write_text(codigo, encoding="utf-8")
        ruta_bitacora = Path(carpeta) / "streamlit.log"
        puerto = puerto_libre()
        with open(ruta_bitacora, "w", encoding="utf-8") as bitacora:
            proceso = arrancar_streamlit(ruta, puerto, bitacora)
            try:
                ok, mensaje, captura = _servir_y_auditar(proceso, puerto)
            except Exception as exc:
                ok, mensaje, captura = False, f"Error durante la inspección de Playwright: {exc}", None
            finally:
                detener(proceso)
        servidor = ruta_bitacora.read_text(encoding="utf-8", errors="ignore").strip()
    if not ok and servidor:
        mensaje += f"\n\nTraceback completo del servidor (stderr):\n{servidor[-4000:]}"
    return ok, mensaje, captura


def validar_en_navegador(codigo: str) -> Tuple[bool, str, Optional[str]]:
    """Fase 2: la app servida por `streamlit run` y recorrida con Playwright.

    Corre en un hilo aparte, como en gs3.py, para no chocar con el bucle de
    asyncio cuando se ejecuta desde un notebook.
    """
    with concurrent.futures.ThreadPoolExecutor(max_workers=1) as ejecutor:
        return ejecutor.submit(_validar_en_navegador, codigo).result()


# --- Los agentes ------------------------------------------------------------

PROMPT_EXTRACTOR = (
    "Eres un Arquitecto de Datos e Información de primer nivel.\n"
    "Analiza holísticamente la información cruda del documento. Tu objetivo es estructurar "
    "toda la información relevante, sin perder cifras aisladas, fechas clave ni relaciones contextuales.\n\n"
    "DEBES estructurar tu salida obligatoriamente bajo este formato JSON:\n"
    "{\n"
    "  \"metadata\": {\n"
    "      \"document_title\": \"Título o tema del documento\",\n"
    "      \"entity_or_subject\": \"Empresa, proyecto, cliente o asunto central\",\n"
    "      \"dates\": [{\"context\": \"Fecha de corte / reporte / junta\", \"value\": \"YYYY-MM-DD o texto\"}],\n"
    "      \"general_summary\": \"Resumen conciso del propósito del documento\"\n"
    "  },\n"
    "  \"global_kpis_and_totals\": [\n"
    "      {\n"
    "          \"metric\": \"Nombre del KPI / Total / Porcentaje\",\n"
    "          \"value\": \"Valor exacto (numérico o texto)\",\n"
    "          \"unit\": \"MXN, USD, %, unidades, etc.\",\n"
    "          \"context\": \"Qué representa o contra qué se compara\"\n"
    "      }\n"
    "  ],\n"
    "  \"tables\": [\n"
    "      {\n"
    "          \"table_id\": \"nombre_o_identificador_tabla\",\n"
    "          \"description\": \"Propósito de esta tabla o desglose\",\n"
    "          \"columns\": [\"columna1\", \"columna2\"],\n"
    "          \"records\": [\n"
    "              {\"columna1\": \"valorA\", \"columna2\": 100}\n"
    "          ]\n"
    "      }\n"
    "  ],\n"
    "  \"semantic_relationships\": [\n"
    "      {\n"
    "          \"source\": \"Origen (ej. global_kpis_and_totals.Gran Total)\",\n"
    "          \"relation_type\": \"agregacion | desglose | comparativa\",\n"
    "          \"target\": \"Destino (ej. tables[0].records.Monto)\",\n"
    "          \"explanation\": \"Explicación de cómo interactúan estos datos\"\n"
    "      }\n"
    "  ]\n"
    "}\n\n"
    "REGLAS:\n"
    "1. No descartes subtotales, márgenes ni porcentajes; colócalos en 'global_kpis_and_totals'.\n"
    "2. Si hay múltiples pestañas o secciones, refléjalas como tablas separadas en la lista 'tables'.\n"
    "3. Incluye una muestra exhaustiva de registros (mínimo 10-25 registros realistas por tabla).\n"
    "4. Devuelve ÚNICAMENTE el bloque de código ```json ... ``` sin comentarios ni introducciones."
)

PROMPT_UX = (
    "Eres un Diseñador Senior de Soluciones de Business Intelligence y Dashboards en Streamlit.\n"
    "Analiza la estructura JSON adjunta y genera las directrices arquitectónicas para crear un "
    "dashboard que muestre ABSOLUTAMENTE TODA la información relevante.\n\n"
    "Proporciona instrucciones claras y detalladas en lenguaje natural sobre:\n"
    "1. Estructura visual superior: Título, metadatos principales del documento y tarjetas `st.metric` "
    "   en columnas para todos los elementos de 'global_kpis_and_totals'.\n"
    "2. Pestañas o secciones con `st.tabs`: Organiza una pestaña para cada tabla presente en el JSON.\n"
    "3. Barra lateral (`st.sidebar`): Filtros dinámicos interactivos (selectores multivariables, "
    "   controles deslizantes de fechas o montos).\n"
    "4. Visualizaciones gráficas: Recomienda gráficos de Plotly (`plotly.express`) interactivos "
    "   (barras, líneas, pie o dispersión) con nombres exactos de variables basados en las relaciones semánticas.\n"
    "5. Visores de datos crudos: Incorpora `st.dataframe` con formato optimizado para explorar los registros.\n\n"
)

PROMPT_DESARROLLADOR = (
    "Eres un Programador Senior en Python especializado en Streamlit y Data Science.\n"
    "Escribe el código fuente completo, autocontenido, funcional y sin errores para `app.py`.\n\n"
    "REGLAS TÉCNICAS OBLIGATORIAS:\n"
    "1. Configuración de página: La PRIMERA llamada de Streamlit debe ser `st.set_page_config(page_title=..., layout='wide')`.\n"
    "2. Autonomía de datos: Construye directamente los DataFrames en memoria (`pd.DataFrame(...)`) con los registros "
    "   completos proporcionados en el JSON para que el script funcione inmediatamente sin archivos externos.\n"
    "3. File Uploader opcional: Agrega en el sidebar un `st.file_uploader(['xlsx', 'csv'])` para cargar nuevos datos; "
    "   si no se sube ninguno, utiliza automáticamente los DataFrames en memoria.\n"
    "4. Estilos y Tablas: Para evitar fallos con dependencias de renderizado, utiliza `st.dataframe(df, width=\"stretch\")`.\n"
    "5. Gráficos Plotly y Leyendas: Usa `st.plotly_chart(fig, width=\"stretch\")`; el parámetro antiguo de ancho "
    "   de contenedor está deprecado en Streamlit, no lo uses.\n"
    "   REGLA ESTRICTA DE SINTAXIS PLOTLY: NUNCA utilices parámetros inexistentes como `legend_position` en `fig.update_layout()`. "
    "   Para posicionar la leyenda horizontalmente abajo, utiliza SIEMPRE el diccionario estándar:\n"
    "   `fig.update_layout(legend=dict(orientation='h', yanchor='bottom', y=-0.25, xanchor='center', x=0.5))`.\n"
    "6. Tipado Defensivo y Filtros Reactivos: Asegura la existencia de columnas (`if 'col' in df.columns:`) y convierte "
    "   columnas numéricas con `pd.to_numeric(df['col'], errors='coerce')`. SIEMPRE valida DataFrames vacíos tras aplicar "
    "   filtros interactivos antes de operar o graficar (`if df_filtrado.empty: st.warning('...'); st.stop()`) para blindar el runtime.\n"
    "7. Prohibido usar rutas de Colab o locales ('/content/...', 'C:\\...').\n"
    "8. Salida: Entrega ÚNICAMENTE el código Python completo encerrado en un bloque ```python ... ``` sin introducciones ni despedidas.\n\n"
)

INSTRUCCIONES_DE_REPARACION = (
    "INSTRUCCIONES DE AUTO-REPARACIÓN:\n"
    "1. Inspecciona minuciosamente el fragmento marcado con '--> [LÍNEA DEL FALLO]' y el traceback del servidor.\n"
    "2. Corrige puntualmente la excepción detectada sin alterar partes funcionales ni introducir regresiones.\n"
    "3. Si se adjunta una captura visual, revisa las tarjetas rojas de error o gráficos vacíos para corregir la UI.\n"
    "4. Declara explícitamente todas las dependencias necesarias e impórtalas al inicio.\n"
    "5. Convierte tipos de datos de forma defensiva (`pd.to_numeric(..., errors='coerce')`) antes de operar o graficar.\n"
    "6. Prohibido invocar métodos, argumentos o propiedades inexistentes en Streamlit o Plotly."
)


@dataclass(frozen=True)
class Modelos:
    extractor: Any
    ux: Any
    desarrollador: Any


@dataclass(frozen=True)
class Validadores:
    simular: Callable[[str], Tuple[bool, str]] = simular_con_apptest
    navegador: Callable[[str], Tuple[bool, str, Optional[str]]] = validar_en_navegador


def agente_extractor(state: GraphState, llm: Any) -> Dict[str, Any]:
    """Agente 1: comprensión del documento y modelado en JSON."""
    print(" [Agente 1] Extrayendo metadatos, métricas globales, tablas y relaciones...")
    contenido = leer_archivo(state["file_path"])
    humano = (HumanMessage(content=contenido) if isinstance(contenido, list)
              else HumanMessage(content=f"CONTENIDO DEL DOCUMENTO:\n{contenido}"))
    respuesta = llm.invoke([SystemMessage(content=PROMPT_EXTRACTOR), humano])
    return {
        "extracted_json": limpiar_bloque_de_codigo(texto_de_respuesta(respuesta), "json"),
        "retry_count": 0, "validation_error": None, "error_history": [],
        "ui_screenshot": None, "is_valid": False,
    }


def agente_ux(state: GraphState, llm: Any) -> Dict[str, Any]:
    """Agente 2: especificación UX/UI del dashboard."""
    print(" [Agente 2] Diseñando wireframe y especificación UX/UI del Dashboard...")
    prompt = PROMPT_UX + f"ESQUEMA DE DATOS:\n{state['extracted_json']}"
    return {"ux_prompt": texto_de_respuesta(llm.invoke([HumanMessage(content=prompt)]))}


def _seccion_de_retroalimentacion(state: GraphState) -> str:
    historial = state.get("error_history") or []
    if not historial:
        return ""
    intentos = "\n\n".join(f"--- INTENTO FALLIDO {i + 1} ---\n{error}" for i, error in enumerate(historial))
    seccion = (f"\n\n HISTORIAL CRONOLÓGICO DE ERRORES Y TRAZABILIDAD:\n{intentos}\n\n"
               f"{INSTRUCCIONES_DE_REPARACION}")
    anterior = state.get("python_code")
    if anterior:
        seccion += ("\n\nCÓDIGO ANTERIOR QUE GENERÓ EL FALLO (APLICA EL PARCHE SOBRE ESTA BASE):\n"
                    f"```python\n{anterior}\n```")
    return seccion


def _mensaje_para_desarrollador(prompt: str, captura: Optional[str]) -> HumanMessage:
    if not captura:
        return HumanMessage(content=prompt)
    texto = (f"{prompt}\n\nCAPTURA VISUAL DEL DASHBOARD:\n"
             "Se adjunta la captura visual en tiempo real donde se manifestó el fallo en Playwright. "
             "Revisa las tarjetas rojas de excepción, contenedores vacíos o desalineaciones visuales "
             "para repararlos puntualmente.")
    return HumanMessage(content=[
        {"type": "text", "text": texto},
        {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{captura}"}},
    ])


def agente_desarrollador(state: GraphState, llm: Any) -> Dict[str, Any]:
    """Agente 3: el `app.py`, con el historial de fallos y la captura si los hay."""
    intento = state.get("retry_count", 0)
    print(f" [Agente 3] Generando app.py autocontenido (Intento {intento + 1})...")
    prompt = (PROMPT_DESARROLLADOR + f"ESQUEMA Y DATOS:\n{state['extracted_json']}\n\n"
              f"ESPECIFICACIÓN UX:\n{state['ux_prompt']}" + _seccion_de_retroalimentacion(state))
    respuesta = llm.invoke([_mensaje_para_desarrollador(prompt, state.get("ui_screenshot"))])
    codigo = sanitizar_codigo(limpiar_bloque_de_codigo(texto_de_respuesta(respuesta), "python"))
    return {"python_code": codigo, "retry_count": intento + 1}


def _fallo(codigo: str, mensaje: str, fase: str, historial: List[str],
           captura: Optional[str]) -> Dict[str, Any]:
    completo = mensaje + contexto_del_fallo(codigo, mensaje)
    print(f"  -> Anomalía detectada ({fase}):\n{completo}")
    historial.append(f"[{fase}]:\n{completo}")
    return {"is_valid": False, "validation_error": completo, "error_history": historial, "ui_screenshot": captura}


def nodo_validador(state: GraphState, validadores: Validadores) -> Dict[str, Any]:
    """Control de calidad: sintaxis + AppTest, y luego la app vista en navegador."""
    print(" [Control de Calidad] Validando sintaxis, dependencias, ciclo reactivo y vista de interfaz...")
    codigo = state["python_code"]
    historial = list(state.get("error_history") or [])
    ok, mensaje = validar_codigo(codigo, validadores.simular)
    if not ok:
        return _fallo(codigo, mensaje, "Fase 1 - AST / AppTest", historial, None)
    print("  -> Inspeccionando vista de interfaz en navegador headless Playwright...")
    ok, mensaje, captura = validadores.navegador(codigo)
    if not ok:
        return _fallo(codigo, mensaje, "Fase 2 - Renderizado Playwright UI", historial, captura)
    print("  -> Verificación integral exitosa. Generando requirements.txt sincronizado...")
    return {"is_valid": True, "validation_error": None, "error_history": historial,
            "ui_screenshot": None, "requirements_txt": extraer_dependencias(codigo)}


def decidir_siguiente(state: GraphState) -> str:
    """Router: publica, reintenta o aborta."""
    if state["is_valid"]:
        return "GitOps_Deploy"
    if state["retry_count"] >= MAX_INTENTOS:
        print(f"  -> Límite de reintentos alcanzado ({MAX_INTENTOS}) con errores no resueltos. Cancelando despliegue.")
        return "Pipeline_Abort"
    return "Code_Generator"


def _escribir_atomico(ruta: Path, texto: str) -> None:
    temporal = ruta.with_name(ruta.name + ".tmp")
    temporal.write_text(texto, encoding="utf-8")
    os.replace(temporal, ruta)


def nodo_publicar(state: GraphState, salida: Path) -> Dict[str, Any]:
    """Deja app.py y requirements.txt en la carpeta del repo; el commit lo hace el Action."""
    print(f" [GitOps Deploy] Escribiendo app.py y requirements.txt en {salida}...")
    salida.mkdir(parents=True, exist_ok=True)
    codigo = state["python_code"] if state["python_code"].endswith("\n") else state["python_code"] + "\n"
    requisitos = state.get("requirements_txt") or extraer_dependencias(codigo)
    _escribir_atomico(salida / "requirements.txt", requisitos)
    _escribir_atomico(salida / "app.py", codigo)
    return {"commit_status": {
        "status_code": ESTADO_PUBLICADO,
        "message": "app.py y requirements.txt listos para el commit del Action.",
        "url": None, "installed_dependencies": requisitos.splitlines(),
    }}


def nodo_abortar(state: GraphState) -> Dict[str, Any]:
    """No publica nada: la liga sigue enseñando el último dashboard que pasó."""
    historial = state.get("error_history") or [state.get("validation_error") or "Error no registrado."]
    return {"commit_status": {
        "status_code": ESTADO_ABORTADO,
        "message": (f"Despliegue cancelado por seguridad: El código no superó las pruebas tras "
                    f"{state['retry_count']} intentos.\n\nHistorial completo de fallos:\n" + "\n\n".join(historial)),
        "url": None, "installed_dependencies": [],
    }}


def construir_grafo(modelos: Modelos, salida: Path, validadores: Optional[Validadores] = None) -> Any:
    validadores = validadores or Validadores()
    flujo = StateGraph(GraphState)
    flujo.add_node("Extractor", functools.partial(agente_extractor, llm=modelos.extractor))
    flujo.add_node("UX_Architect", functools.partial(agente_ux, llm=modelos.ux))
    flujo.add_node("Code_Generator", functools.partial(agente_desarrollador, llm=modelos.desarrollador))
    flujo.add_node("Code_Validator", functools.partial(nodo_validador, validadores=validadores))
    flujo.add_node("GitOps_Deploy", functools.partial(nodo_publicar, salida=salida))
    flujo.add_node("Pipeline_Abort", nodo_abortar)
    flujo.add_edge(START, "Extractor")
    flujo.add_edge("Extractor", "UX_Architect")
    flujo.add_edge("UX_Architect", "Code_Generator")
    flujo.add_edge("Code_Generator", "Code_Validator")
    flujo.add_conditional_edges("Code_Validator", decidir_siguiente, {
        "Code_Generator": "Code_Generator", "GitOps_Deploy": "GitOps_Deploy", "Pipeline_Abort": "Pipeline_Abort",
    })
    flujo.add_edge("GitOps_Deploy", END)
    flujo.add_edge("Pipeline_Abort", END)
    return flujo.compile()


def ejecutar(ruta: str, modelos: Modelos, salida: Path, validadores: Optional[Validadores] = None) -> Dict[str, Any]:
    estado_inicial: GraphState = {"file_path": ruta, "retry_count": 0, "validation_error": None,
                                  "error_history": [], "ui_screenshot": None, "is_valid": False}
    return construir_grafo(modelos, salida, validadores).invoke(estado_inicial)


def construir_modelos(entorno: Optional[Mapping[str, str]] = None) -> Modelos:
    """Los tres Gemini, con una clave por agente o `GEMINI_API_KEY` para todos."""
    entorno = os.environ if entorno is None else entorno
    claves, faltan = [], []
    for numero in (1, 2, 3):
        clave = (entorno.get(f"GEMINI_API_KEY_AGENT_{numero}") or entorno.get("GEMINI_API_KEY") or "").strip()
        claves.append(clave)
        if not clave:
            faltan.append(f"GEMINI_API_KEY_AGENT_{numero}")
    if faltan:
        raise ConfiguracionIncompleta(
            "Faltan las claves de Gemini: " + ", ".join(faltan) + " (o GEMINI_API_KEY para los tres agentes).")
    from langchain_google_genai import ChatGoogleGenerativeAI

    modelo = (entorno.get("GEMINI_MODEL") or "").strip() or MODELO_POR_DEFECTO
    return Modelos(*(ChatGoogleGenerativeAI(model=modelo, google_api_key=clave) for clave in claves))


# --- Descarga del documento -------------------------------------------------

class _SinRedirecciones(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req: Any, fp: Any, code: int, msg: str, headers: Any, newurl: str) -> None:
        raise OrigenNoPermitido(f"El documento redirige a otra dirección (HTTP {code}); no se sigue.")


def nombre_permitido(url: str, origen: str) -> str:
    """El nombre seguro del documento, si la URL cae dentro de `origen`."""
    if not origen:
        raise ConfiguracionIncompleta(
            "ORIGEN_PERMITIDO está vacío: debe ser la URL pública del Storage de Holtmont, "
            "p. ej. https://<proyecto>.supabase.co/storage/v1/object/public/")
    permitido = urllib.parse.urlsplit(origen if origen.endswith("/") else origen + "/")
    pedido = urllib.parse.urlsplit(url)
    ruta = posixpath.normpath(urllib.parse.unquote(pedido.path))
    if (pedido.scheme, pedido.netloc) != (permitido.scheme, permitido.netloc) or not ruta.startswith(permitido.path):
        raise OrigenNoPermitido(f"La URL {url} no está dentro de {origen}.")
    nombre = re.sub(r"[^A-Za-z0-9._-]", "_", posixpath.basename(ruta))
    extension = nombre.rsplit(".", 1)[-1].lower() if "." in nombre else ""
    if extension not in EXTENSIONES_SOPORTADAS:
        raise ArchivoIlegible(f"El formato .{extension} no se puede convertir en dashboard.")
    return nombre


def descargar_archivo(url: str, destino: Path, origen: str, max_bytes: int = MAX_BYTES_ARCHIVO,
                      abridor: Optional[urllib.request.OpenerDirector] = None) -> Path:
    """Baja el documento del Storage de Holtmont, con tope de tamaño."""
    nombre = nombre_permitido(url, origen)
    abridor = abridor or urllib.request.build_opener(_SinRedirecciones)
    try:
        with abridor.open(url, timeout=30) as respuesta:
            datos = respuesta.read(max_bytes + 1)
    except urllib.error.HTTPError as exc:
        raise ArchivoIlegible(f"Storage respondió {exc.code} al descargar {nombre}.") from exc
    except urllib.error.URLError as exc:
        raise ArchivoIlegible(f"No se pudo descargar {nombre}: {exc.reason}") from exc
    if len(datos) > max_bytes:
        raise ArchivoIlegible(f"{nombre} pesa más de {max_bytes} bytes.")
    ruta = destino / nombre
    ruta.write_bytes(datos)
    return ruta


# --- Línea de comandos ------------------------------------------------------

def _argumentos(argv: Optional[List[str]]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Genera el dashboard de la Pre Work Order.")
    fuente = parser.add_mutually_exclusive_group(required=True)
    fuente.add_argument("--url", help="URL pública del documento en el Storage de Holtmont")
    fuente.add_argument("--archivo", help="Ruta local del documento (uso manual)")
    parser.add_argument("--salida", default=".", help="Carpeta donde se escriben app.py y requirements.txt")
    parser.add_argument("--folio", default="", help="Folio de la Pre Work Order, para la bitácora")
    return parser.parse_args(argv)


def _imprimir_reporte(reporte: Dict[str, Any], folio: str) -> None:
    print("\n" + "=" * 60 + "\nREPORTE FINAL DE GITOPS\n" + "=" * 60)
    if folio:
        print(f"Folio: {folio}")
    print(f"Estado: {reporte['status_code']}")
    print(f"Mensaje: {reporte['message']}")
    print(f"Dependencias autogeneradas: {reporte.get('installed_dependencies')}")


def main(argv: Optional[List[str]] = None, construir: Callable[[], Modelos] = construir_modelos,
         validadores: Optional[Validadores] = None) -> int:
    """0 si app.py quedó listo para el commit; 1 si no (el Action termina en rojo)."""
    args = _argumentos(argv)
    print("Iniciando orquestación de agentes...\n" + "=" * 60)
    try:
        verificar_aislamiento()
        modelos = construir()
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = (Path(args.archivo) if args.archivo
                    else descargar_archivo(args.url, Path(carpeta), os.environ.get("ORIGEN_PERMITIDO", "")))
            estado = ejecutar(str(ruta), modelos, Path(args.salida), validadores)
    except ErrorDelPipeline as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    _imprimir_reporte(estado["commit_status"], args.folio)
    return 0 if estado["commit_status"]["status_code"] == ESTADO_PUBLICADO else 1


if __name__ == "__main__":
    sys.exit(main())
